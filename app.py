# -*- coding: utf-8 -*-
import streamlit as st
import json, io
from openai import OpenAI
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from docx import Document  # 用于导出 Word 教案

# 1. 页面配置
st.set_page_config(page_title="全能 AI 化学教研工作站", page_icon="🧪", layout="wide")

st.title("🧪 全能 AI 高中化学教研工作站")
st.caption("一键自动生成：16:9卡片式PPT课件 + Word标准化教案 + 高考级随堂练习题")

# 侧边栏配置
st.sidebar.header("⚙️ API 与参数设置")
api_key = st.sidebar.text_input("DeepSeek API Key", value="sk-45096723d48b4687a215456ae97f30b3", type="password")
base_url = st.sidebar.text_input("Base URL", value="https://api.deepseek.com")

# 主界面输入
col_input1, col_input2 = st.columns([2, 1])
with col_input1:
    topic = st.text_input("📌 化学课题名称：", "氧化还原反应")
with col_input2:
    grade = st.selectbox("🎯 适用年级/模块：", [
        "高中化学必修一", 
        "高中化学必修二", 
        "选择性必修一（化学反应原理）", 
        "选择性必修三（有机化学）"
    ])

# --- 核心逻辑函数 ---

# 1. 绘制 PPT
def build_ppt(slides_data):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    
    COLOR_BG = RGBColor(245, 247, 250)
    COLOR_PRIMARY = RGBColor(16, 37, 66)
    COLOR_ACCENT = RGBColor(230, 81, 0)
    COLOR_CARD = RGBColor(255, 255, 255)
    COLOR_TEXT = RGBColor(51, 51, 51)
    
    for page in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = COLOR_BG
        
        # 标题
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.8))
        p = tb.text_frame.paragraphs[0]
        p.text = page["title"]
        p.font.size, p.font.bold, p.font.color.rgb = Pt(28), True, COLOR_PRIMARY
        
        # 装饰线
        line = slide.shapes.add_shape(1, Inches(0.8), Inches(1.3), Inches(11.73), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb, line.line.color.rgb = COLOR_ACCENT, COLOR_ACCENT
        
        # 内容卡片
        card = slide.shapes.add_shape(1, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = RGBColor(220, 224, 230)
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left, tf.margin_top = Inches(0.4), Inches(0.4)
        
        for j, pt in enumerate(page.get("points", [])):
            p_pt = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            p_pt.text = f"•  {pt}"
            p_pt.font.size = Pt(20)
            p_pt.font.color.rgb = COLOR_TEXT
            p_pt.space_after = Pt(16)
            
    ppt_out = io.BytesIO()
    prs.save(ppt_out)
    ppt_out.seek(0)
    return ppt_out

# 2. 生成 Word 教案（已修正变量名统一为 topic 和 grade）
def build_doc(topic, grade, slides_data, quiz_data):
    doc = Document()
    doc.add_heading(f'《{topic}》教学设计教案', 0)
    doc.add_paragraph(f'适用年级：{grade}')
    
    doc.add_heading('一、教学大纲与板书要点', level=1)
    for s in slides_data:
        doc.add_heading(s['title'], level=2)
        for p in s.get('points', []):
            doc.add_paragraph(p, style='List Bullet')
            
    doc.add_heading('二、配套随堂检测题', level=1)
    for idx, q in enumerate(quiz_data, 1):
        doc.add_paragraph(f"{idx}. {q['question']}")
        for opt in q.get('options', []):
            doc.add_paragraph(opt)
        p_ans = doc.add_paragraph()
        p_ans.add_run(f"【答案】{q['answer']}\n【解析】{q['explanation']}").bold = True
        doc.add_paragraph('')
        
    doc_out = io.BytesIO()
    doc.save(doc_out)
    doc_out.seek(0)
    return doc_out

# --- 按钮与执行 ---
if st.button("🚀 一键生成全套教研资源（PPT + 教案 + 题库）", type="primary", use_container_width=True):
    if not api_key:
        st.error("请先在左侧输入你的 API Key！")
    else:
        with st.spinner("🤖 AI 正在智能思考并构建教学全套方案，请稍候..."):
            try:
                client = OpenAI(api_key=api_key, base_url=base_url)
                
                # 调大模型同时生成 PPT + 试题
                prompt = f"""
                你是一位高中化学特级教师。请为【{grade}】课题【{topic}】设计教研资源。
                请严格输出为合法 JSON 格式（不要带任何 ```json 标记），格式如下：
                {{
                  "slides": [
                    {{"title": "页面标题", "points": ["要点1", "要点2"]}}
                  ],
                  "quiz": [
                    {{
                      "question": "题目描述",
                      "options": ["A. xx", "B. xx", "C. xx", "D. xx"],
                      "answer": "A",
                      "explanation": "解析过程..."
                    }}
                  ]
                }}
                要求：slides 生成 4 页；quiz 生成 2 道高质量选择题。
                """
                
                response = client.chat.completions.create(
                    model="deepseek-chat",
                    messages=[
                        {"role": "system", "content": "你是一个只输出 JSON 数据的化学专家。"},
                        {"role": "user", "content": prompt}
                    ]
                )
                
                ai_content = response.choices[0].message.content.strip()
                if ai_content.startswith("```"):
                    ai_content = ai_content.split("\n", 1)[1].rsplit("\n", 1)[0]
                
                res_data = json.loads(ai_content)
                slides_data = res_data["slides"]
                quiz_data = res_data["quiz"]
                
                # 生成二进制文件（传参完全一致）
                ppt_file = build_ppt(slides_data)
                doc_file = build_doc(topic, grade, slides_data, quiz_data)
                
                st.success("🎉 全套教研资源生成完毕！")
                
                # 标签页展示
                tab1, tab2, tab3 = st.tabs(["📊 PPT 课件大纲", "✍️ 随堂检测题", "📥 一键导出下载"])
                
                with tab1:
                    for s in slides_data:
                        with st.expander(f"📖 {s['title']}"):
                            for p in s.get("points", []):
                                st.write(f"- {p}")
                                
                with tab2:
                    for idx, q in enumerate(quiz_data, 1):
                        st.markdown(f"**{idx}. {q['question']}**")
                        for opt in q.get("options", []):
                            st.write(opt)
                        st.info(f"**答案**：{q['answer']} \n\n **解析**：{q['explanation']}")
                        st.divider()
                        
                with tab3:
                    col_dl1, col_dl2 = st.columns(2)
                    with col_dl1:
                        st.download_button(
                            label="📥 下载 PPT 课件 (.pptx)",
                            data=ppt_file,
                            file_name=f"{grade}_{topic}_教学课件.pptx",
                            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                            use_container_width=True
                        )
                    with col_dl2:
                        st.download_button(
                            label="📄 下载 Word 教案与试题 (.docx)",
                            data=doc_file,
                            file_name=f"{grade}_{topic}_教案试题.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )
            except Exception as e:
                st.error(f"⚠️ 发生错误：{e}")